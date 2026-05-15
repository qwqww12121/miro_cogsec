from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


DOCS = Path(__file__).resolve().parent
ASSETS = DOCS / "ppt_assets"
OUT = DOCS / "COGSEC_COMPETITION_SUPPORT_DECK.pptx"


def rgb(hex_code: str) -> RGBColor:
    hex_code = hex_code.lstrip("#")
    return RGBColor(int(hex_code[0:2], 16), int(hex_code[2:4], 16), int(hex_code[4:6], 16))


BG = rgb("#F8FAFC")
WHITE = rgb("#FFFFFF")
TEXT = rgb("#0F172A")
SUBTLE = rgb("#334155")
MUTED = rgb("#64748B")
BLUE = rgb("#0C4A6E")
BLUE_BG = rgb("#EFF6FF")
BLUE_LINE = rgb("#BAE6FD")
GREEN_BG = rgb("#F0FDF4")
GREEN_LINE = rgb("#BBF7D0")
ORANGE_BG = rgb("#FFF7ED")
ORANGE_LINE = rgb("#FED7AA")
PURPLE_BG = rgb("#FAF5FF")
PURPLE_LINE = rgb("#DDD6FE")
SOFT_BG = rgb("#F8FAFC")
SOFT_LINE = rgb("#E2E8F0")


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_full_bg(slide, color):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_textbox(
    slide,
    text,
    left,
    top,
    width,
    height,
    *,
    font_name="Microsoft YaHei",
    font_size=20,
    color=TEXT,
    bold=False,
    fill=None,
    line=None,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.08,
):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.clear()
    tf.text = text
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color
        p.alignment = align
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def add_title(slide, text):
    return add_textbox(
        slide,
        text,
        0.35,
        0.18,
        12.2,
        0.5,
        font_name="Microsoft YaHei UI",
        font_size=24,
        color=TEXT,
        bold=True,
    )


def add_source(slide, text):
    return add_textbox(
        slide,
        f"资料来源：{text}",
        0.35,
        7.0,
        12.2,
        0.24,
        font_size=8.5,
        color=MUTED,
    )


def add_card(slide, title, body, left, top, width, height, fill, line):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    add_textbox(
        slide,
        title,
        left + 0.08,
        top + 0.05,
        width - 0.16,
        0.35,
        font_name="Microsoft YaHei UI",
        font_size=17,
        color=TEXT,
        bold=True,
    )
    add_textbox(
        slide,
        body,
        left + 0.08,
        top + 0.38,
        width - 0.16,
        height - 0.45,
        font_size=14,
        color=SUBTLE,
    )
    return shape


def add_picture(slide, path, left, top, width, height, caption=None):
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    if caption:
        add_textbox(slide, caption, left, top + height + 0.03, width, 0.24, font_size=8.5, color=MUTED)


def add_code_panel(slide, title, code, left, top, width, height):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = SOFT_BG
    shape.line.color.rgb = SOFT_LINE
    add_textbox(
        slide,
        title,
        left + 0.08,
        top + 0.05,
        width - 0.16,
        0.28,
        font_name="Microsoft YaHei UI",
        font_size=14,
        color=TEXT,
        bold=True,
    )
    add_textbox(
        slide,
        code,
        left + 0.08,
        top + 0.34,
        width - 0.16,
        height - 0.42,
        font_name="Consolas",
        font_size=10.5,
        color=TEXT,
    )


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, BG)
add_textbox(slide, "MiroFish / CogSec", 0.5, 0.9, 6.5, 0.6, font_name="Microsoft YaHei UI", font_size=28, color=TEXT, bold=True)
add_textbox(slide, "基于反事实推演的个人决策图谱", 0.5, 1.55, 7.2, 0.6, font_name="Microsoft YaHei UI", font_size=22, color=BLUE, bold=True)
add_textbox(
    slide,
    "配套报告：COGSEC_COMPETITION_SUMMARY_REPOSITIONED.tex\n版本：v3.2\n日期：2026-04-22",
    0.5,
    2.45,
    4.1,
    1.0,
    font_size=17,
    color=SUBTLE,
)
add_textbox(
    slide,
    "面向诈骗、钓鱼和社会工程场景，在用户执行高风险动作前给出路径比较、可逆性窗口和干预建议。",
    0.5,
    3.85,
    6.2,
    1.0,
    font_size=18,
    color=TEXT,
    fill=BLUE_BG,
    line=BLUE_LINE,
)
add_picture(slide, ASSETS / "mirofish_logo.jpeg", 8.1, 1.25, 4.0, 2.6, "项目 logo / 本地开源仓库资产")
add_source(slide, "本地素材：static/image/MiroFish_logo_compressed.jpeg；报告正文：docs/COGSEC_COMPETITION_SUMMARY_REPOSITIONED.tex")


# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "1. 这个系统是做什么的")
add_textbox(
    slide,
    "它不是只判断像不像诈骗，而是帮助用户在危险动作发生前做决定：\n\n• 识别当前场景中的高危动作与诱导线索\n• 比较继续执行与暂停核验两条路径\n• 输出风险差异、可逆性窗口和个性化干预建议",
    0.45,
    1.0,
    5.3,
    3.2,
    font_size=19,
    color=TEXT,
)
add_card(slide, "输入", "文本 / 聊天记录 / 截图 / 场景描述", 6.0, 1.25, 1.75, 1.05, BLUE_BG, BLUE_LINE)
add_card(slide, "识别", "T0 护栏 + 风险因子 + ThreatKnowledgeRAG", 7.95, 1.25, 1.75, 1.05, GREEN_BG, GREEN_LINE)
add_card(slide, "Fork", "Branch A / Branch B 反事实推演", 9.9, 1.25, 1.95, 1.05, ORANGE_BG, ORANGE_LINE)
add_card(slide, "输出", "风险差异 / 可逆性 / 干预处方", 7.95, 3.05, 2.05, 1.05, PURPLE_BG, PURPLE_LINE)
add_textbox(
    slide,
    "一句话：把模糊风险转换成可比较、可解释、可干预的决策图谱。",
    6.0,
    4.9,
    5.85,
    0.8,
    font_name="Microsoft YaHei UI",
    font_size=20,
    color=TEXT,
    bold=True,
    fill=SOFT_BG,
    line=SOFT_LINE,
)
add_source(slide, "报告正文：执行摘要 / 它是做什么的")


# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "2. 为什么基于 MiroFish 可以构建个人决策图谱")
add_textbox(
    slide,
    "MiroFish 原始强项不是某个单场景功能，而是复杂情境建模闭环：\n\n• 非结构化输入 -> 图谱/状态对象\n• 多角色、多因素约束下的环境构造\n• 状态演化与关键分叉跟踪\n• 报告输出与交互解释\n\n这正好可以从世界排演缩放到个体关键动作前的局部平行路径。",
    0.45,
    0.95,
    4.8,
    4.6,
    font_size=18,
    color=TEXT,
)
add_picture(slide, ASSETS / "run_1.png", 5.75, 0.95, 6.2, 4.3, "本地运行截图：原始 MiroFish 首页与输入入口")
add_textbox(
    slide,
    "方法迁移逻辑：平行世界 -> 局部平行路径；群体演化 -> 个体决策分叉。",
    0.45,
    5.55,
    11.5,
    0.7,
    font_name="Microsoft YaHei UI",
    font_size=19,
    color=BLUE,
    bold=True,
    fill=BLUE_BG,
    line=BLUE_LINE,
)
add_source(slide, "MiroFish GitHub：https://github.com/666ghj/MiroFish；MiroFish-Offline：https://github.com/nikmcfly/MiroFish-Offline")


# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "3. 原版 MiroFish 与当前项目的关键差异")
add_card(
    slide,
    "原版 MiroFish",
    "• 群体/环境的未来排演\n• 输入是报告、新闻、政策、故事等现实种子\n• 输出是预测报告与交互式数字世界\n• 重心是图谱构建、环境搭建与大规模 Agent 模拟",
    0.45,
    1.0,
    5.2,
    3.85,
    SOFT_BG,
    SOFT_LINE,
)
add_card(
    slide,
    "当前 CogSec 项目",
    "• 个体在诈骗场景中的关键决策路径\n• 输入是对话、文本、截图、问卷与上下文\n• 输出是个人决策图谱、双分支对照、可逆性与处方\n• 重心是 T0/T1、风险因子、RAG、Fork 与 Reporter",
    6.1,
    1.0,
    5.75,
    3.85,
    BLUE_BG,
    BLUE_LINE,
)
add_textbox(
    slide,
    "准确说法：以 MiroFish 为情境建模底座，以 CogSec 为认知安全主线。",
    0.45,
    5.25,
    11.4,
    0.75,
    font_name="Microsoft YaHei UI",
    font_size=20,
    color=TEXT,
    bold=True,
    fill=ORANGE_BG,
    line=ORANGE_LINE,
)
add_source(slide, "对照来源：本地 README、README-EN 与报告正文 为什么基于 MiroFish 可以构建个人决策图谱")


# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "4. 同类竞品与开源调研")
add_card(slide, "消费者 AI 反诈", "Bitdefender Scamio\nNorton Genie\n\n更像内容层第二意见，强在快筛与低门槛。", 0.45, 1.0, 2.45, 2.15, BLUE_BG, BLUE_LINE)
add_card(slide, "企业人因风险平台", "SoSafe\nKnowBe4 AIDA\nAdaptive Security\n\n重培训闭环与组织长期风险下降。", 3.0, 1.0, 2.45, 2.15, GREEN_BG, GREEN_LINE)
add_card(slide, "SOC / Triage Agent", "Microsoft Defender\nPhishing Triage Agent\n\n服务分析师工作流，而非普通用户的即时决策辅助。", 5.55, 1.0, 2.45, 2.15, ORANGE_BG, ORANGE_LINE)
add_card(slide, "训练 / 开源 / 研究", "Pretexta\nAXIS\n\n前者偏训练与复盘，后者偏多智能体反事实解释研究。", 8.1, 1.0, 3.0, 2.15, PURPLE_BG, PURPLE_LINE)
add_textbox(
    slide,
    "结论：已存在大量检测、训练、triage 与研究解释方案，但个人关键动作前的反事实决策图谱仍有明显空位。",
    0.45,
    3.6,
    11.4,
    1.0,
    font_name="Microsoft YaHei UI",
    font_size=20,
    color=TEXT,
    bold=True,
    fill=SOFT_BG,
    line=SOFT_LINE,
)
add_textbox(
    slide,
    "代表链接：\nBitdefender Scamio / Norton Genie / SoSafe / KnowBe4 AIDA / Adaptive Security / Microsoft Defender / Pretexta / AXIS",
    0.45,
    5.0,
    11.4,
    0.9,
    font_size=15,
    color=SUBTLE,
)
add_source(slide, "产品与开源入口见 PPT_MATERIALS.md；报告正文 同类竞品与 GitHub 开源调研")


# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "5. 为什么系统要这样设计：理论与方法依据")
add_card(slide, "双系统理论", "高压力、高不确定性下，System 1 更容易接管决策。\n=> 支持 T0 / T1 双层结构。", 0.45, 1.0, 2.45, 2.15, BLUE_BG, BLUE_LINE)
add_card(slide, "用户情境与个体差异", "同样内容对不同人、不同上下文的风险并不一样。\n=> 支持 persona_state_vector。", 3.0, 1.0, 2.45, 2.15, GREEN_BG, GREEN_LINE)
add_card(slide, "说服原则", "权威、紧迫、社会认同、承诺一致性、互惠、喜好。\n=> 支持攻击-人-环境图谱。", 5.55, 1.0, 2.45, 2.15, ORANGE_BG, ORANGE_LINE)
add_card(slide, "反事实解释", "不只说明发生了什么，更说明换一条路径会怎样。\n=> 支持 Branch A / Branch B。", 8.1, 1.0, 3.0, 2.15, PURPLE_BG, PURPLE_LINE)
add_textbox(
    slide,
    "因此，T0 + 风险向量 + ThreatKnowledgeRAG + Fork + Reporter 不是堆模块，而是从问题本质、竞品边界与理论依据共同推导出的组合。",
    0.45,
    3.6,
    11.4,
    1.1,
    font_name="Microsoft YaHei UI",
    font_size=19,
    color=TEXT,
    bold=True,
    fill=SOFT_BG,
    line=SOFT_LINE,
)
add_source(slide, "Frontiers 2025 / NIST Phish Scale / PMC social engineering cognition / persuasion survey / AXIS")


# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "6. 系统总体架构与主链工作流")
add_card(slide, "输入治理层", "文本 / 聊天记录 / 截图 OCR / 问卷 / 上下文 + 隐私脱敏", 0.45, 1.1, 2.4, 1.25, BLUE_BG, BLUE_LINE)
add_card(slide, "分析理解层", "风险因子提取、场景识别、ThreatKnowledgeRAG、攻击-人-环境图谱", 3.05, 1.1, 2.4, 1.25, GREEN_BG, GREEN_LINE)
add_card(slide, "反事实推演层", "WorldState、Fork、Branch A / B、风险差值与可逆性", 5.65, 1.1, 2.4, 1.25, ORANGE_BG, ORANGE_LINE)
add_card(slide, "报告与干预层", "数字孪生摘要、路径对照、干预窗口、个性化处方", 8.25, 1.1, 3.05, 1.25, PURPLE_BG, PURPLE_LINE)
add_textbox(
    slide,
    "Input -> PrivacySanitizer -> PersonaStateVector -> RiskGraphBundle -> WorldState -> Fork A/B -> Counterfactual Risk -> Intervention",
    0.45,
    3.0,
    11.4,
    0.7,
    font_name="Consolas",
    font_size=16,
    color=TEXT,
    fill=SOFT_BG,
    line=SOFT_LINE,
)
add_textbox(
    slide,
    "T0：验证码 / 转账 / 屏幕共享 / 安全账户等红旗的毫秒级护栏\nT1：画像、图谱、Fork 比较和反事实报告生成",
    0.45,
    4.15,
    11.4,
    1.0,
    font_size=19,
    color=SUBTLE,
)
add_source(slide, "本地 README、docs/COGSEC_MIROFISH_MAINLINE_REFACTOR.md、报告正文 系统总体架构")


# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "7. 现有界面与产品形态证据")
add_picture(slide, ASSETS / "run_3.png", 0.35, 0.85, 7.65, 4.25, "本地运行截图：图谱可视化 + 环境搭建 / 双栏工作流")
add_picture(slide, ASSETS / "run_6.png", 8.2, 0.85, 4.65, 4.25, "本地运行截图：复杂图谱与节点详情")
add_textbox(
    slide,
    "前端当前最有价值的证据不是概念稿，而是已经存在的工作流与工作台结构：\n\n• 首页 / 主流程 / 报告页 / 交互页\n• CogSecWorkbench 五屏：用户数字孪生、攻击-人-环境图谱、双分支反事实推演、风险 / 可逆性曲线、个性化干预处方",
    0.45,
    5.45,
    11.4,
    1.1,
    font_size=18,
    color=TEXT,
)
add_source(slide, "本地截图：static/image/Screenshot/运行截图3.png、运行截图6.png；组件：frontend/src/components/cogsec/CogSecWorkbench.vue")


# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "8. 关键代码证据")
code1 = """self.profile_extractor = CognitiveProfileExtractor(...)\nself.threat_rag = ThreatKnowledgeRAG(Config.CHROMA_PATH)\nself.runtime = MiroFishRuntime()\nself.reporter = CounterfactualReporter()\nself.t0_responder = T0FastResponder(...)\nself.privacy_sanitizer = PrivacySanitizer(...)\n\nruntime_result = self.runtime.run(...)\nreport = self.reporter.generate_report(...)"""
code2 = """@app.get('/health')\n@app.get('/api/cogsec/sample')\n@app.post('/api/cogsec/analyze')\n\nconst screens = [\n  { key: 'twin', label: '1. 用户数字孪生' },\n  { key: 'graph', label: '2. 攻击-人-环境图谱' },\n  { key: 'fork', label: '3. 双分支反事实推演' },\n  { key: 'curve', label: '4. 风险 / 可逆性曲线' },\n  { key: 'prescription', label: '5. 个性化干预处方' }\n]"""
add_code_panel(slide, "后端主链编排：backend/app/services/cogsec_service.py", code1, 0.45, 1.0, 5.65, 3.55)
add_code_panel(slide, "最小 Demo 与五屏工作台：run_cogsec_demo.py / CogSecWorkbench.vue", code2, 6.25, 1.0, 5.6, 3.55)
add_textbox(
    slide,
    "结论：项目已经不是 PPT 先行 的概念稿，后端主链对象、最小 Demo 入口与五屏工作台都已有明确代码落点。",
    0.45,
    5.1,
    11.4,
    0.9,
    font_name="Microsoft YaHei UI",
    font_size=19,
    color=TEXT,
    bold=True,
    fill=SOFT_BG,
    line=SOFT_LINE,
)
add_source(slide, "backend/app/services/cogsec_service.py；backend/run_cogsec_demo.py；frontend/src/components/cogsec/CogSecWorkbench.vue")


# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "9. 创新点与差异化价值")
add_card(slide, "从内容识别转向决策建模", "最小分析单位不是消息文本，而是用户下一步的高危动作。", 0.45, 1.0, 2.45, 2.0, BLUE_BG, BLUE_LINE)
add_card(slide, "个体风险因子进入路径推演", "风险向量用于解释为什么同样刺激对不同人作用不同。", 3.0, 1.0, 2.45, 2.0, GREEN_BG, GREEN_LINE)
add_card(slide, "反事实报告作为主输出", "系统主打的是 Branch A / B 差异，而不是单一风险分。", 5.55, 1.0, 2.45, 2.0, ORANGE_BG, ORANGE_LINE)
add_card(slide, "T0 / T1 双层防御结构", "把毫秒级护栏与深度解释结合起来，兼顾阻断与说明。", 8.1, 1.0, 3.0, 2.0, PURPLE_BG, PURPLE_LINE)
add_textbox(
    slide,
    "创新点成立的前提：\n• 页面里真的看得到关键节点、双分支和可逆性\n• 风险因子服务于解释，而不是夸张的人格诊断\n• 用样例、测试和 benchmark 证明链路稳定",
    0.45,
    3.45,
    11.4,
    1.5,
    font_size=19,
    color=TEXT,
)
add_source(slide, "报告正文 创新点与差异化价值")


# Slide 11
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, WHITE)
add_title(slide, "10. 风险治理与迭代路线")
add_card(
    slide,
    "治理边界",
    "• 仅用于防御与决策辅助\n• 优先本地/前置脱敏\n• 个体风险因子不形成惩罚性标签\n• 证据不足时允许返回需要人工核验",
    0.45,
    1.0,
    5.2,
    2.55,
    SOFT_BG,
    SOFT_LINE,
)
add_card(
    slide,
    "比赛版近期优先级",
    "• 稳定个人决策 DAG 可视化\n• 强化 Branch A / Branch B 报告与可逆性曲线\n• 补标准场景样例与截图\n• 用 benchmark/测试结果支撑性能表述\n• 统一首页与答辩主叙事口径",
    6.1,
    1.0,
    5.75,
    2.55,
    BLUE_BG,
    BLUE_LINE,
)
add_textbox(
    slide,
    "中长期可扩展：更丰富的风险因子动态更新、更强多模态输入、更完整知识库与标注体系，以及复杂关系链/长期演化场景下的多智能体环境建模。",
    0.45,
    4.1,
    11.4,
    1.0,
    font_size=18,
    color=TEXT,
)
add_source(slide, "报告正文 风险治理、隐私与伦理边界；当前局限与迭代路线")


# Slide 12
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_full_bg(slide, BG)
add_title(slide, "11. 资料入口与结束页")
add_textbox(
    slide,
    "配套文件：\n• COGSEC_COMPETITION_SUMMARY_REPOSITIONED.tex\n• PPT_MATERIALS.md\n• ppt_assets/\n\n关键公开链接：\n• MiroFish  https://github.com/666ghj/MiroFish\n• Pretexta  https://github.com/dalpan/Pretexta\n• AXIS  https://arxiv.org/abs/2505.17801\n• NIST Phish Scale  https://www.nist.gov/publications/phishing-user-context-understanding-nist-phish-scale",
    0.45,
    1.1,
    7.9,
    4.0,
    font_size=17,
    color=TEXT,
)
add_textbox(
    slide,
    "一句话收束：在损失真正发生前，为个体提供一张可解释、可比较、可干预的决策图谱。",
    0.45,
    5.4,
    7.9,
    0.95,
    font_name="Microsoft YaHei UI",
    font_size=21,
    color=TEXT,
    bold=True,
    fill=WHITE,
    line=SOFT_LINE,
)
add_picture(slide, ASSETS / "mirofish_logo.jpeg", 9.0, 1.6, 2.7, 1.8, "项目本地素材")
add_source(slide, "更多资料见 PPT_MATERIALS.md")


prs.save(str(OUT))
print(f"Generated: {OUT}")
